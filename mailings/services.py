from __future__ import annotations

from collections.abc import Callable, Iterable, Iterator, Sequence
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.core.validators import validate_email
from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from .models import ImportEntry, ImportRun, MailingRecord

REQUIRED_HEADERS = ("external_id", "user_id", "email", "subject", "message")


class ImportFileError(ValueError):
    """Raised when a workbook cannot be imported at all."""


class RowValidationError(ValueError):
    """Raised for one malformed workbook row."""


@dataclass(frozen=True, slots=True)
class MailingCandidate:
    row_number: int
    external_id: str
    user_id: int
    email: str
    subject: str
    message: str


@dataclass(slots=True)
class ImportStats:
    processed: int = 0
    created: int = 0
    skipped: int = 0
    errors: int = 0
    sent: int = 0
    delivery_failed: int = 0


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _required_text(value: Any, field: str, max_length: int | None = None) -> str:
    if _is_blank(value):
        raise RowValidationError(f"{field} is required")
    text = str(value).strip()
    if max_length is not None and len(text) > max_length:
        raise RowValidationError(f"{field} exceeds {max_length} characters")
    return text


def _external_id(value: Any) -> str:
    if isinstance(value, bool):
        raise RowValidationError("external_id must be a string or number")
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return _required_text(value, "external_id", 128)


def _user_id(value: Any) -> int:
    if isinstance(value, bool):
        raise RowValidationError("user_id must be a positive integer")
    try:
        number = int(value)
    except (TypeError, ValueError, OverflowError) as exc:
        raise RowValidationError("user_id must be a positive integer") from exc
    if isinstance(value, float) and not value.is_integer():
        raise RowValidationError("user_id must be a positive integer")
    if not 1 <= number <= 2**31 - 1:
        raise RowValidationError("user_id must be between 1 and 2147483647")
    return number


def _candidate(row_number: int, values: dict[str, Any]) -> MailingCandidate:
    email = _required_text(values["email"], "email", 254)
    try:
        validate_email(email)
    except ValidationError as exc:
        raise RowValidationError("email is invalid") from exc
    return MailingCandidate(
        row_number=row_number,
        external_id=_external_id(values["external_id"]),
        user_id=_user_id(values["user_id"]),
        email=email,
        subject=_required_text(values["subject"], "subject", 255),
        message=_required_text(values["message"], "message"),
    )


def _header_indexes(worksheet: Worksheet) -> dict[str, int]:
    header_row = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if header_row is None:
        raise ImportFileError("workbook is empty")
    indexes: dict[str, int] = {}
    for index, raw_name in enumerate(header_row):
        if _is_blank(raw_name):
            continue
        name = str(raw_name).strip().lower()
        if name in indexes:
            raise ImportFileError(f"duplicate header: {name}")
        indexes[name] = index
    missing = [name for name in REQUIRED_HEADERS if name not in indexes]
    if missing:
        raise ImportFileError(f"missing required headers: {', '.join(missing)}")
    return indexes


def _rows(worksheet: Worksheet, indexes: dict[str, int]) -> Iterator[tuple[int, dict[str, Any]]]:
    for row_number, row in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
        if all(_is_blank(value) for value in row):
            continue
        yield row_number, {name: row[index] if index < len(row) else None for name, index in indexes.items()}


def _batches(source: Iterable[Any], size: int) -> Iterator[list[Any]]:
    batch: list[Any] = []
    for item in source:
        batch.append(item)
        if len(batch) == size:
            yield batch
            batch = []
    if batch:
        yield batch


def _persist_batch(
    rows: Sequence[tuple[int, dict[str, Any]]],
    import_run: ImportRun,
    stats: ImportStats,
    on_row_error: Callable[[int, str], None],
) -> None:
    batch = ImportStats(processed=len(rows))
    candidates: dict[str, MailingCandidate] = {}
    row_errors: list[tuple[int, str]] = []
    for row_number, values in rows:
        try:
            item = _candidate(row_number, values)
        except RowValidationError as exc:
            batch.errors += 1
            row_errors.append((row_number, str(exc)))
        else:
            if item.external_id in candidates:
                batch.skipped += 1
            else:
                candidates[item.external_id] = item

    # Commit jobs, memberships and counters together. A crash loses at most the
    # unfinished batch; re-importing safely reconstructs its file membership.
    with transaction.atomic():
        existing = dict(MailingRecord.objects.filter(external_id__in=candidates).values_list("external_id", "pk"))
        fresh = [item for key, item in candidates.items() if key not in existing]
        batch.skipped += len(existing)
        # Hold user rows only for this short batch, protecting the FK against a
        # concurrent deletion. All writers acquire user/record locks in ID order.
        user_ids = set(
            get_user_model()
            .objects.select_for_update()
            .filter(pk__in={item.user_id for item in fresh})
            .order_by("pk")
            .values_list("pk", flat=True)
        )
        valid = []
        for item in fresh:
            if item.user_id not in user_ids:
                batch.errors += 1
                row_errors.append((item.row_number, f"user_id {item.user_id} does not exist"))
            else:
                valid.append(item)
        records = [
            MailingRecord(
                external_id=item.external_id,
                user_id=item.user_id,
                email=item.email,
                subject=item.subject,
                message=item.message,
                import_run=import_run,
            )
            for item in sorted(valid, key=lambda item: item.external_id)
        ]
        MailingRecord.objects.bulk_create(records, batch_size=len(records) or 1, ignore_conflicts=True)
        # Read rows back to distinguish our inserts from a competing import's.
        resolved = list(
            MailingRecord.objects.filter(
                external_id__in=(*existing, *(item.external_id for item in valid))
            ).values_list("pk", "import_run_id", "external_id")
        )
        batch.created = sum(
            owner == import_run.pk and external_id not in existing for _, owner, external_id in resolved
        )
        batch.skipped += len(records) - batch.created
        ImportEntry.objects.bulk_create(
            [ImportEntry(import_run=import_run, record_id=pk) for pk, _, _ in sorted(resolved)],
            ignore_conflicts=True,
        )
        for name in ("processed", "created", "skipped", "errors"):
            field = "error_count" if name == "errors" else f"{name}_count"
            setattr(import_run, field, getattr(stats, name) + getattr(batch, name))
        import_run.save(update_fields=("processed_count", "created_count", "skipped_count", "error_count"))
    for name in ("processed", "created", "skipped", "errors"):
        setattr(stats, name, getattr(stats, name) + getattr(batch, name))
    for row_number, error in row_errors:
        on_row_error(row_number, error)


def import_xlsx(
    source: str | Path,
    *,
    batch_size: int = 500,
    on_row_error: Callable[[int, str], None] | None = None,
) -> tuple[ImportRun, ImportStats]:
    path = Path(source)
    if path.suffix.lower() != ".xlsx":
        raise ImportFileError("source file must have the .xlsx extension")
    if not path.is_file():
        raise ImportFileError(f"file does not exist: {path}")
    if not 1 <= batch_size <= 1000:
        raise ImportFileError("batch size must be between 1 and 1000")

    report_error = on_row_error or (lambda _row, _message: None)
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise ImportFileError(f"cannot read workbook: {exc}") from exc
    import_run: ImportRun | None = None
    stats = ImportStats()
    try:
        worksheet = workbook.active
        indexes = _header_indexes(worksheet)
        import_run = ImportRun.objects.create(source_filename=path.name)
        for batch in _batches(_rows(worksheet, indexes), batch_size):
            _persist_batch(batch, import_run, stats, report_error)
        _finish_import(import_run, ImportRun.Status.IMPORTED)
        return import_run, stats
    except KeyboardInterrupt:
        if import_run is not None:
            _finish_import(import_run, ImportRun.Status.INTERRUPTED)
        raise
    except Exception:
        if import_run is not None:
            _finish_import(import_run, ImportRun.Status.FAILED)
        raise
    finally:
        workbook.close()


def _finish_import(import_run: ImportRun, status: str) -> None:
    import_run.status = status
    import_run.finished_at = timezone.now()
    import_run.save(update_fields=("status", "finished_at"))
